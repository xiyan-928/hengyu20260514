"""
自动控制脚本引擎与 xlsx 加载单元测试。
"""
import os
import sys
import tempfile
import unittest
from unittest.mock import MagicMock, patch

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

from auto_control_context import AutoControlContext
from auto_control_manager import AutoControlManager
from auto_control_script_engine import AutoControlScriptEngine, get_script_engine
from Devices.device_settings import DeviceSettings
from script_define_loader import (
    ScriptRegistry,
    load_script_registry,
    resolve_script_define_path,
)


def _reset_engine() -> AutoControlScriptEngine:
    e = get_script_engine()
    e.set_registry(ScriptRegistry())
    e._active_hook = None
    e._steps = []
    e._applied_count = 0
    e._next_deadline = None
    e._on_complete = None
    return e


def _build_minimal_xlsx(path: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    w1 = wb.create_sheet("基础点位配置")
    w1.append(["名称", "点位"])
    w1.append(["V1", 0])
    w1.append(["V2", 1])
    w2 = wb.create_sheet("模式定义")
    w2.append(["NAME", "V1", "V2", "备注", "Visible"])
    w2.append(["V1", 1, 0, "", 1])
    w2.append(["V2", 0, 1, "", 1])
    w3 = wb.create_sheet("函数定义")
    w3.append(["函数名", "函数定义"])
    w3.append(["run_on_alarm_on", "V1=1 V2=2"])
    w3.append(["run_on_alarm_off", "V2=3 V1=4"])
    wb.save(path)
    wb.close()


class TestScriptDefineLoader(unittest.TestCase):
    def test_load_minimal_xlsx(self):
        with tempfile.TemporaryDirectory() as d:
            p = os.path.join(d, "s.xlsx")
            _build_minimal_xlsx(p)
            reg = load_script_registry(p)
            self.assertEqual(reg.points.get("V1"), 0)
            self.assertEqual(reg.points.get("V2"), 1)
            self.assertIn("V1", reg.mode_segments)
            self.assertIn("run_on_alarm_on", reg.hook_steps)
            self.assertIn("run_on_alarm_off", reg.hook_steps)
            steps = reg.hook_steps["run_on_alarm_on"]
            self.assertEqual(steps, [("V1", 1.0), ("V2", 2.0)])
            self.assertEqual(reg.hook_raw_commands["run_on_alarm_on"], "V1=1 V2=2")
            self.assertEqual(reg.hook_steps["run_on_alarm_off"], [("V2", 3.0), ("V1", 4.0)])

    def test_resolve_path_relative_to_hy_online(self):
        hy = _ROOT
        out = resolve_script_define_path("define/script_define.xlsx", hy_online_root=hy)
        self.assertTrue(out.endswith(os.path.join("define", "script_define.xlsx")))


class TestAutoControlScriptEngine(unittest.TestCase):
    def tearDown(self):
        _reset_engine()

    def test_start_writes_first_mode_and_tick_advances(self):
        e = _reset_engine()
        reg = ScriptRegistry()
        reg.points = {"A": 0}
        reg.mode_segments = {"M1": [(0, [True])], "M2": [(0, [False])]}
        reg.hook_steps = {"run_on_alarm_on": [("M1", 0.01), ("M2", 0.01)]}
        reg.hook_raw_commands = {"run_on_alarm_on": "M1=0.01 M2=0.01"}
        e.set_registry(reg)

        client = MagicMock()
        rr = MagicMock()
        rr.isError.return_value = False
        client.write_coils.return_value = rr
        ctx = AutoControlContext(
            settings=DeviceSettings.instance(),
            coils_modbus_client=client,
            coils_modbus_host="192.168.1.10",
            coils_modbus_port=502,
            modbus_unit_id=1,
        )
        clock = [0.0]

        def fake_mono():
            return clock[0]

        with patch("auto_control_script_engine.time.monotonic", side_effect=fake_mono):
            e.start("run_on_alarm_on", ctx)
        client.write_coils.assert_called_with(0, [True], unit=1)
        self.assertEqual(ctx.current_hook_name, "run_on_alarm_on")
        self.assertEqual(ctx.current_mode_name, "M1")
        self.assertEqual(ctx.current_script_command, "M1=0.01 M2=0.01")
        client.reset_mock()

        with patch("auto_control_script_engine.time.monotonic", side_effect=fake_mono):
            clock[0] = 0.02
            e.tick(ctx)
        client.write_coils.assert_called_with(0, [False], unit=1)
        client.reset_mock()

        with patch("auto_control_script_engine.time.monotonic", side_effect=fake_mono):
            clock[0] = 0.03
            e.tick(ctx)
        self.assertIsNone(e._active_hook)

    def test_start_empty_hook_calls_on_complete(self):
        e = _reset_engine()
        reg = ScriptRegistry()
        reg.hook_steps = {}
        e.set_registry(reg)
        done = MagicMock()
        ctx = AutoControlContext(settings=DeviceSettings.instance(), modbus_client=None)
        e.start("run_on_warmup", ctx, on_complete=done)
        done.assert_called_once()

    def test_interrupt_restarts_script(self):
        e = _reset_engine()
        reg = ScriptRegistry()
        reg.mode_segments = {"M1": [(0, [True])]}
        reg.hook_steps = {"h1": [("M1", 100.0)], "h2": [("M1", 100.0)]}
        e.set_registry(reg)
        client = MagicMock()
        rr = MagicMock()
        rr.isError.return_value = False
        client.write_coils.return_value = rr
        ctx = AutoControlContext(
            settings=DeviceSettings.instance(),
            modbus_client=client,
            modbus_unit_id=1,
        )
        e.start("h1", ctx)
        e.start("h2", ctx)
        self.assertEqual(e._active_hook, "h2")

    @patch("auto_control_script_engine.run_ephemeral")
    @patch("auto_control_script_engine.app_config.is_mock_mode", return_value=False)
    def test_write_mode_uses_modbus_coils_endpoint_when_no_coils_client(self, _mock_is_mock, mock_run_ephemeral):
        e = _reset_engine()
        reg = ScriptRegistry()
        reg.mode_segments = {"M1": [(0, [True])]}
        e.set_registry(reg)

        fake_client = MagicMock()
        rr = MagicMock()
        rr.isError.return_value = False
        fake_client.write_coils.return_value = rr

        def invoke_task(ip, port, task, context=None):
            task(fake_client)

        mock_run_ephemeral.side_effect = invoke_task

        ctx = AutoControlContext(
            settings=DeviceSettings.instance(),
            modbus_client=None,
            modbus_host="localhost",
            modbus_port=5020,
            coils_modbus_host="192.168.1.10",
            coils_modbus_port=502,
            modbus_unit_id=1,
        )

        e.write_mode("M1", ctx)

        mock_run_ephemeral.assert_called_once()
        args, kwargs = mock_run_ephemeral.call_args
        self.assertEqual(args[0], "192.168.1.10")
        self.assertEqual(args[1], 502)
        self.assertEqual(kwargs["context"]["mode_name"], "M1")
        fake_client.write_coils.assert_called_with(0, [True], unit=1)

    def test_write_mode_prefers_explicit_coils_client(self):
        e = _reset_engine()
        reg = ScriptRegistry()
        reg.mode_segments = {"M1": [(0, [True])]}
        e.set_registry(reg)

        monitor_client = MagicMock()
        coils_client = MagicMock()
        rr = MagicMock()
        rr.isError.return_value = False
        coils_client.write_coils.return_value = rr

        ctx = AutoControlContext(
            settings=DeviceSettings.instance(),
            modbus_client=monitor_client,
            modbus_host="localhost",
            modbus_port=5020,
            coils_modbus_client=coils_client,
            coils_modbus_host="192.168.1.10",
            coils_modbus_port=502,
            modbus_unit_id=1,
        )

        e.write_mode("M1", ctx)

        monitor_client.write_coils.assert_not_called()
        coils_client.write_coils.assert_called_with(0, [True], unit=1)

    def test_http_apply_mode_source_uses_modbus_coils_endpoint(self):
        src = open(os.path.join(_ROOT, "main.py"), "r", encoding="utf-8").read()

        self.assertIn("host, port = settings.get_modbus_coils_endpoint()", src)
        self.assertIn('detail=f"无法连接线圈 Modbus {host}:{port}"', src)


class TestWarmupSingleStart(unittest.TestCase):
    """预热：同一 needs_warmup 会话内 run_on_warmup 只触发一次"""

    @patch("auto_control_manager.run_on_warmup")
    def test_second_warmup_phase_does_not_call_again(self, mock_run):
        mgr = AutoControlManager(lambda: None)
        ctx = mgr._build_context(None, {"unit_id": 1, "monitor_register": 1})
        mgr._has_warmup = False
        mgr._needs_warmup = True
        mgr._warmup_phase(ctx)
        mgr._warmup_phase(ctx)
        mock_run.assert_called_once()


if __name__ == "__main__":
    unittest.main()
