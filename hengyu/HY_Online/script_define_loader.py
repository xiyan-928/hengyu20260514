"""
从 HY_Online/define/script_define.xlsx 加载自动控制脚本：基础点位、模式、函数定义。

表头必须与模板 ``HY_Online/define/script_define.xlsx`` 一致（列名见本模块常量），
配置项 auto_control.script_define_path 默认为相对 HY_Online 根的 ``define/script_define.xlsx``
（由 resolve_script_define_path 展开）。
"""
from __future__ import annotations

import logging
import os
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

logger = logging.getLogger(__name__)

# Sheet 名称（与 Excel 模板一致）
SHEET_POINTS = "基础点位配置"
SHEET_MODES = "模式定义"
SHEET_FUNCTIONS = "函数定义"

# 表头：与 define/script_define.xlsx 模板首行一致（匹配时去空白、英文列大小写不敏感）
# 基础点位配置：名称 | 点位
POINT_NAME_HEADERS = ("名称",)
POINT_ADDR_HEADERS = ("点位",)

# 模式定义：NAME | <与基础点位「名称」同名的各列> | 备注 | Visible
MODE_NAME_HEADERS = ("NAME",)
MODE_VISIBLE_HEADERS = ("Visible",)
MODE_REMARK_HEADERS = ("备注",)

# 函数定义：函数名 | 函数定义
FUNC_NAME_HEADERS = ("函数名",)
FUNC_DEF_HEADERS = ("函数定义",)


def _norm_header(s: str) -> str:
    return str(s).strip().lower()


def _cell_str(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _cell_int(v) -> Optional[int]:
    if v is None or v == "":
        return None
    try:
        return int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None


def _split_contiguous_writes(
    addr_to_bool: Dict[int, bool],
) -> List[Tuple[int, List[bool]]]:
    """按地址升序拆成若干段连续线圈，每段一次 write_coils。"""
    if not addr_to_bool:
        return []
    addrs = sorted(addr_to_bool.keys())
    segments: List[Tuple[int, List[bool]]] = []
    seg_start = addrs[0]
    seg_vals = [addr_to_bool[addrs[0]]]
    for i in range(1, len(addrs)):
        a = addrs[i]
        if a == addrs[i - 1] + 1:
            seg_vals.append(addr_to_bool[a])
        else:
            segments.append((seg_start, seg_vals))
            seg_start = a
            seg_vals = [addr_to_bool[a]]
    segments.append((seg_start, seg_vals))
    return segments


@dataclass
class ScriptRegistry:
    """内存中的脚本定义。"""

    # 点位名称 -> 线圈地址
    points: Dict[str, int] = field(default_factory=dict)
    # 模式名 -> [(起始地址, [bool,...]), ...]
    mode_segments: Dict[str, List[Tuple[int, List[bool]]]] = field(default_factory=dict)
    # 模式定义表行序：(模式名, 是否在 UI 显示按钮, 按钮文案)
    mode_ui_rows: List[Tuple[str, bool, str]] = field(default_factory=list)
    # hook 名 -> [(模式名, 该步执行后的等待秒数), ...]
    hook_steps: Dict[str, List[Tuple[str, float]]] = field(default_factory=dict)
    # hook 名 -> 原始函数定义文本
    hook_raw_commands: Dict[str, str] = field(default_factory=dict)
    load_errors: List[str] = field(default_factory=list)

    def is_empty(self) -> bool:
        return not self.points and not self.mode_segments and not self.hook_steps

    def list_visible_mode_buttons(self) -> List[Dict[str, str]]:
        """Visible 为真的模式，顺序与表格行一致。"""
        return [
            {"mode": name, "label": label}
            for name, visible, label in self.mode_ui_rows
            if visible
        ]


_STEP_RE = re.compile(r"(\w+)\s*=\s*(\d+(?:\.\d+)?)")


def _parse_function_definition(text: str) -> Optional[List[Tuple[str, float]]]:
    """解析如 V1=240 V2=2 -> [('V1', 240.0), ('V2', 2.0)]"""
    if not text or not str(text).strip():
        return None
    s = str(text).strip()
    out: List[Tuple[str, float]] = []
    for m in _STEP_RE.finditer(s):
        mode = m.group(1)
        sec = float(m.group(2))
        out.append((mode, sec))
    if not out:
        logger.warning(f"函数定义无法解析出任何 模式=秒 步骤: {text!r}")
        return None
    return out


def _parse_visible_cell(raw) -> bool:
    """缺省为 True；0/false/否 等为 False。"""
    if raw is None or str(raw).strip() == "":
        return True
    s = str(raw).strip().lower()
    if s in ("0", "false", "no", "n", "否", "off"):
        return False
    if s in ("1", "true", "yes", "y", "是", "on"):
        return True
    try:
        return bool(int(float(s)))
    except (TypeError, ValueError):
        return True


def _find_col(headers: List[str], candidates: Tuple[str, ...]) -> Optional[int]:
    nh = [_norm_header(h) for h in headers]
    for c in candidates:
        cn = _norm_header(c)
        for i, h in enumerate(nh):
            if h == cn:
                return i
    return None


def _load_points_sheet(ws) -> Tuple[Dict[str, int], List[str]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}, ["基础点位配置: 空表"]
    header = [str(x).strip() if x is not None else "" for x in rows[0]]
    ic_name = _find_col(header, POINT_NAME_HEADERS)
    ic_addr = _find_col(header, POINT_ADDR_HEADERS)
    errs: List[str] = []
    if ic_name is None or ic_addr is None:
        return {}, [
            "基础点位配置: 首行须为模板列「名称」「点位」（与 define/script_define.xlsx 一致）"
        ]
    points: Dict[str, int] = {}
    for ri, row in enumerate(rows[1:], start=2):
        name = _cell_str(row[ic_name] if ic_name < len(row) else None)
        addr = _cell_int(row[ic_addr] if ic_addr < len(row) else None)
        if not name:
            continue
        if addr is None:
            errs.append(f"基础点位配置 行{ri}: 名称 {name!r} 地址无效")
            continue
        if name in points:
            errs.append(f"基础点位配置 行{ri}: 重复名称 {name!r}")
        points[name] = addr
    return points, errs


def _load_modes_sheet(
    ws, points: Dict[str, int]
) -> Tuple[Dict[str, List[Tuple[int, List[bool]]]], List[str], List[Tuple[str, bool, str]]]:
    rows = list(ws.iter_rows(values_only=True))
    errs: List[str] = []
    mode_ui_rows: List[Tuple[str, bool, str]] = []
    if not rows:
        return {}, ["模式定义: 空表"], []
    header = [str(x).strip() if x is not None else "" for x in rows[0]]
    ic_mode = _find_col(header, MODE_NAME_HEADERS)
    if ic_mode is None:
        return (
            {},
            [
                "模式定义: 首行须包含模板列 NAME（与 define/script_define.xlsx 一致）"
            ],
            [],
        )
    ic_visible = _find_col(header, MODE_VISIBLE_HEADERS)
    ic_remark = _find_col(header, MODE_REMARK_HEADERS)
    point_cols: Dict[str, int] = {}
    for j, h in enumerate(header):
        if j == ic_mode:
            continue
        if ic_visible is not None and j == ic_visible:
            continue
        if ic_remark is not None and j == ic_remark:
            continue
        hn = _cell_str(h)
        if hn in points:
            point_cols[hn] = j
    if not point_cols:
        errs.append("模式定义: 未找到与基础点位匹配的列名")
    modes: Dict[str, List[Tuple[int, List[bool]]]] = {}
    for ri, row in enumerate(rows[1:], start=2):
        mode_name = _cell_str(row[ic_mode] if ic_mode < len(row) else None)
        if not mode_name:
            continue
        addr_to_val: Dict[int, bool] = {}
        for pname, col_idx in point_cols.items():
            if col_idx >= len(row):
                continue
            raw = row[col_idx]
            if raw is None or str(raw).strip() == "":
                continue
            try:
                v = int(float(str(raw).strip()))
            except (TypeError, ValueError):
                errs.append(f"模式定义 行{ri} 列{pname}: 非 0/1 数值")
                continue
            addr_to_val[points[pname]] = bool(v)
        if not addr_to_val:
            errs.append(f"模式定义 行{ri}: 模式 {mode_name!r} 无有效点位值")
            continue
        segs = _split_contiguous_writes(addr_to_val)
        modes[mode_name] = segs
        vis = _parse_visible_cell(
            row[ic_visible] if ic_visible is not None and ic_visible < len(row) else None
        )
        remark = _cell_str(
            row[ic_remark] if ic_remark is not None and ic_remark < len(row) else None
        )
        label = remark if remark else mode_name
        mode_ui_rows.append((mode_name, vis, label))
    return modes, errs, mode_ui_rows


def _load_functions_sheet(
    ws, valid_modes: set
) -> Tuple[Dict[str, List[Tuple[str, float]]], Dict[str, str], List[str]]:
    rows = list(ws.iter_rows(values_only=True))
    errs: List[str] = []
    if not rows:
        return {}, {}, ["函数定义: 空表"]
    header = [str(x).strip() if x is not None else "" for x in rows[0]]
    ic_name = _find_col(header, FUNC_NAME_HEADERS)
    ic_def = _find_col(header, FUNC_DEF_HEADERS)
    if ic_name is None or ic_def is None:
        return {}, {}, [
            "函数定义: 首行须为模板列「函数名」「函数定义」（与 define/script_define.xlsx 一致）"
        ]
    hooks: Dict[str, List[Tuple[str, float]]] = {}
    raw_commands: Dict[str, str] = {}
    for ri, row in enumerate(rows[1:], start=2):
        hname = _cell_str(row[ic_name] if ic_name < len(row) else None)
        def_text = row[ic_def] if ic_def < len(row) else None
        def_text_clean = _cell_str(def_text)
        if not hname:
            continue
        steps = _parse_function_definition(def_text_clean)
        if not steps:
            errs.append(f"函数定义 行{ri}: {hname!r} 定义无效或为空")
            continue
        for mode, _ in steps:
            if mode not in valid_modes:
                errs.append(
                    f"函数定义 行{ri}: hook {hname!r} 引用未知模式 {mode!r}"
                )
        if any(mode not in valid_modes for mode, _ in steps):
            continue
        hooks[hname] = steps
        raw_commands[hname] = def_text_clean
    return hooks, raw_commands, errs


def load_script_registry(path: str) -> ScriptRegistry:
    """
    从 xlsx 加载。文件不存在或解析失败时返回带 load_errors 的 registry，不抛异常。
    """
    reg = ScriptRegistry()
    if not path or not os.path.isfile(path):
        msg = f"script_define 文件不存在或不可读: {path}"
        logger.warning(msg)
        reg.load_errors.append(msg)
        return reg
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        msg = f"未安装 openpyxl，无法加载脚本定义: {e}"
        logger.error(msg)
        reg.load_errors.append(msg)
        return reg
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        msg = f"打开 script_define 失败: {path}: {e}"
        logger.error(msg)
        reg.load_errors.append(msg)
        return reg
    try:
        def get_sheet(title: str):
            if title in wb.sheetnames:
                return wb[title]
            return None

        ws_p = get_sheet(SHEET_POINTS)
        if ws_p is None and wb.sheetnames:
            ws_p = wb[wb.sheetnames[0]]
            reg.load_errors.append(f"未找到工作表 {SHEET_POINTS!r}，使用第一个表作为基础点位")
        if ws_p is None:
            reg.load_errors.append("工作簿无工作表")
            return reg

        points, e1 = _load_points_sheet(ws_p)
        reg.points = points
        reg.load_errors.extend(e1)

        ws_m = get_sheet(SHEET_MODES)
        if ws_m is None and len(wb.sheetnames) > 1:
            ws_m = wb[wb.sheetnames[1]]
            reg.load_errors.append(f"未找到工作表 {SHEET_MODES!r}，使用第二个表作为模式定义")
        mode_segments: Dict[str, List[Tuple[int, List[bool]]]] = {}
        if ws_m is not None:
            mode_segments, e2, mode_ui_rows = _load_modes_sheet(ws_m, points)
            reg.load_errors.extend(e2)
            reg.mode_ui_rows = mode_ui_rows
        reg.mode_segments = mode_segments

        valid_modes = set(mode_segments.keys())

        ws_f = get_sheet(SHEET_FUNCTIONS)
        if ws_f is None and len(wb.sheetnames) > 2:
            ws_f = wb[wb.sheetnames[2]]
            reg.load_errors.append(f"未找到工作表 {SHEET_FUNCTIONS!r}，使用第三个表作为函数定义")
        hook_steps: Dict[str, List[Tuple[str, float]]] = {}
        if ws_f is not None:
            hook_steps, raw_commands, e3 = _load_functions_sheet(ws_f, valid_modes)
            reg.load_errors.extend(e3)
            reg.hook_raw_commands = raw_commands
        reg.hook_steps = hook_steps

        logger.debug(
            f"已加载 script_define: 点位={len(reg.points)} 模式={len(reg.mode_segments)} 函数={len(reg.hook_steps)}"
        )
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return reg


def resolve_script_define_path(path: str, hy_online_root: Optional[str] = None) -> str:
    """相对路径相对 HY_Online 根目录解析。"""
    if not path:
        return ""
    if os.path.isabs(path):
        return path
    root = hy_online_root or os.path.dirname(os.path.abspath(__file__))
    return os.path.normpath(os.path.join(root, path))
