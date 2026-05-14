"""
项目目录、SPC 落盘、染料/浴比 xlsx。
"""

from __future__ import annotations

import logging
import os
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from spectrum_app_config import load_config

logger = logging.getLogger(__name__)

_CATEGORY_PREFIX: Dict[str, str] = {
    "dark": "暗光谱",
    "water": "清水光谱",
    "stock": "打样原液光谱",
    "residual": "打样残液光谱",
}

_DYE_XLSX = "染料信息.xlsx"
_FABRIC_XLSX = "浴比布重.xlsx"

_NAME_RE = re.compile(r"^[^\x00/\\:*?\"<>|]+$")


def validate_project_name(name: str) -> str:
    n = (name or "").strip()
    if not n or ".." in n or "/" in n or "\\" in n:
        raise ValueError("无效的项目名称")
    if not _NAME_RE.match(n):
        raise ValueError("项目名称包含非法字符")
    return n


def get_project_root() -> str:
    cfg = load_config()
    return (cfg.get("project_root") or "").strip()


def resolve_project_root() -> Path:
    root = get_project_root()
    if not root:
        raise ValueError("未设置项目根目录（spectrum_app_config 或 SPECTRUM_PROJECT_ROOT）")
    p = Path(root).resolve()
    if not p.is_dir():
        raise ValueError(f"项目根目录不存在或不是文件夹: {p}")
    return p


def project_path(name: str) -> Path:
    n = validate_project_name(name)
    root = resolve_project_root()
    return root / n


def ensure_project(name: str) -> Path:
    p = project_path(name)
    if not p.is_dir():
        raise ValueError(f"项目不存在: {name}")
    return p


def list_projects() -> List[str]:
    try:
        root = resolve_project_root()
    except ValueError:
        return []
    names: List[str] = []
    for child in sorted(root.iterdir(), key=lambda x: x.name.lower()):
        if child.is_dir() and not child.name.startswith("."):
            names.append(child.name)
    return names


def create_project(name: str) -> Path:
    n = validate_project_name(name)
    root = resolve_project_root()
    target = root / n
    if target.exists():
        raise ValueError(f"项目已存在: {n}")
    target.mkdir(parents=True, exist_ok=False)
    return target


def spc_filename_for_category(category: str, integration_time: int, scans_to_average: int) -> str:
    if category not in _CATEGORY_PREFIX:
        raise ValueError(f"未知光谱类别: {category}")
    label = _CATEGORY_PREFIX[category]
    return f"{label}_it{int(integration_time)}_avg{int(scans_to_average)}.spc"


def save_spc_under_project(
    project_name: str,
    category: str,
    wavelengths: List[float],
    spectrum: List[float],
    integration_time: int,
    scans_to_average: int,
    write_spc_fn,
) -> Dict[str, Any]:
    """
    write_spc_fn: Devices.get_spc.write_spc
    """
    folder = ensure_project(project_name)
    fname = spc_filename_for_category(category, integration_time, scans_to_average)
    out_path = folder / fname
    path = write_spc_fn(wavelengths, spectrum, str(out_path))
    return {
        "file_path": path,
        "file_name": os.path.basename(path),
        "category": category,
    }


def read_project_meta(project_name: str) -> Dict[str, Any]:
    """从 xlsx 读取 meta；文件不存在则返回空结构。"""
    folder = ensure_project(project_name)
    dyes: List[Dict[str, Any]] = []
    fabric_weight_g: Optional[float] = None
    bath_ratio: Optional[float] = None

    dye_path = folder / _DYE_XLSX
    fab_path = folder / _FABRIC_XLSX

    try:
        from openpyxl import load_workbook
    except ImportError:
        logger.warning("openpyxl 未安装，无法读取 xlsx")
        return {
            "dyes": dyes,
            "fabric_weight_g": fabric_weight_g,
            "bath_ratio": bath_ratio,
        }

    if dye_path.is_file():
        wb = load_workbook(dye_path, read_only=True)
        try:
            ws = wb.active
            rows = list(ws.iter_rows(min_row=2, values_only=True))
            for row in rows:
                if row and row[0] is not None:
                    dyes.append(
                        {
                            "name": str(row[0]),
                            "ratio": float(row[1]) if len(row) > 1 and row[1] is not None else 0.0,
                        }
                    )
        finally:
            wb.close()

    if fab_path.is_file():
        wb = load_workbook(fab_path, read_only=True)
        try:
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) >= 2:
                r1 = rows[1]
                if r1 and r1[0] is not None:
                    try:
                        fabric_weight_g = float(r1[0])
                    except (TypeError, ValueError):
                        pass
                if r1 and len(r1) > 1 and r1[1] is not None:
                    try:
                        bath_ratio = float(r1[1])
                    except (TypeError, ValueError):
                        pass
        finally:
            wb.close()

    return {
        "dyes": dyes,
        "fabric_weight_g": fabric_weight_g,
        "bath_ratio": bath_ratio,
    }


def write_project_meta(
    project_name: str,
    dyes: List[Dict[str, Any]],
    fabric_weight_g: float,
    bath_ratio: float,
) -> None:
    from openpyxl import Workbook

    folder = ensure_project(project_name)

    wb = Workbook()
    ws = wb.active
    ws.title = "染料"
    ws.append(["染料名称", "比例"])
    for d in dyes:
        ws.append([d.get("name", ""), d.get("ratio", 0.0)])
    wb.save(folder / _DYE_XLSX)
    wb.close()

    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "浴比布重"
    ws2.append(["布重(g)", "浴比"])
    ws2.append([fabric_weight_g, bath_ratio])
    wb2.save(folder / _FABRIC_XLSX)
    wb2.close()


def _latest_spc_path(folder: Path, category: str) -> Optional[Path]:
    """同一类别下取修改时间最新的 .spc（支持不同 it/avg 多文件）。"""
    if category not in _CATEGORY_PREFIX:
        return None
    prefix = _CATEGORY_PREFIX[category]
    candidates = [
        p
        for p in folder.iterdir()
        if p.is_file() and p.suffix.lower() == ".spc" and p.name.startswith(prefix)
    ]
    if not candidates:
        return None
    return max(candidates, key=lambda p: p.stat().st_mtime)


def parse_it_avg_from_spc_filename(filename: str) -> Tuple[int, int]:
    m = re.search(r"_it(\d+)_avg(\d+)\.spc$", filename, re.IGNORECASE)
    if m:
        return int(m.group(1)), int(m.group(2))
    return 10000, 3


def read_spc_wavelengths_intensities(path: Path) -> Tuple[List[float], List[float]]:
    """读取项目生成的 SPC，并兼容旧版错误布局。"""
    try:
        from Devices.read_spc import read_spc_xy
    except ImportError as e:
        raise RuntimeError(f"无法导入 SPC 读取模块: {e}") from e

    wavelengths, spectrum = read_spc_xy(path)
    wl = [float(x) for x in wavelengths.tolist()]
    sp = [float(y) for y in spectrum.tolist()]
    if len(wl) != len(sp):
        raise RuntimeError("波长与强度长度不一致")
    return wl, sp


def get_saved_spectra_for_project(project_name: str) -> Dict[str, Optional[Dict[str, Any]]]:
    """
    从项目目录加载各类别「最新」已保存 SPC 的曲线数据。
    若某类别无文件或读取失败，对应值为 null。
    """
    folder = ensure_project(project_name)
    out: Dict[str, Optional[Dict[str, Any]]] = {}
    for cat in _CATEGORY_PREFIX:
        path = _latest_spc_path(folder, cat)
        if path is None:
            out[cat] = None
            continue
        try:
            wl, sp = read_spc_wavelengths_intensities(path)
            it, avg = parse_it_avg_from_spc_filename(path.name)
            out[cat] = {
                "wavelengths": wl,
                "spectrum": sp,
                "file_name": path.name,
                "integration_time": it,
                "scans_to_average": avg,
            }
        except Exception as e:
            logger.warning("读取 %s 失败: %s", path, e)
            out[cat] = None
    return out
